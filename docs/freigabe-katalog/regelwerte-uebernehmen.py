"""
E-211: Zurückgeschickte Prüfliste der Regelwerte (scripts/regelwerte-pruefpaket.py) in einem Schritt eintragen: Block im
Register docs/freigabe-katalog/freigaben-regelwerte.yaml (Nachweis ohne Namen, E-197 bis E-199) und `pruefung` der
Regelwert-Versionen im Katalog (wie anwenden.py), danach das Bündel der App.

    python3 docs/freigabe-katalog/regelwerte-uebernehmen.py <rücklauf.xlsx> --art mail --datum JJJJ-MM-TT [--mit-anmerkung id1,id2]
        [--umfang "…"] [--korrekturen <datei.yaml>] [--probe]

- Die Rolle ergibt sich aus der Datei (Titel der Anleitung) und muss zu jedem Regelwert passen (regelwerte.rolle).
- Übernommen werden Zeilen mit Ergebnis „richtig“, deren ID, Version und „Wert + Einheit“ noch dem Katalog entsprechen
  (sonst hat sich der Wert seit dem Versand geändert → neu vorlegen). Zeilen mit Anmerkung nur, wenn ihre ID in
  --mit-anmerkung steht (Anmerkung gelesen, nichts zu ändern).
- „falsch → …“: nichts wird geändert; die Zeilen werden aufgelistet (Wert als neue Version anlegen, erneut vorlegen).
- --art: mail (Bestätigung per Mail; Stand = Datum der Mail), blatt (ausgefüllte Datei) oder bericht. --probe schreibt nichts.
- E-213: --korrekturen <datei.yaml> (außerhalb des Repositorys oder im Scratch; keine Namen, kein Mailtext) für Zeilen, deren
  Anmerkung Vergleich, „gültig ab“, Rechtsgrundlage oder Wert präzisiert bzw. korrigiert, und für „falsch“-Zeilen, die der
  Prüfer in der korrigierten Form ausdrücklich bestätigt hat. Die korrigierte Fassung muss vorher als neue Version im Katalog
  stehen – genau so, wie der Prüfer sie angegeben hat. Je Gruppe ein Block mit eigenem, genauem Umfang:
      - umfang: "Wert, Vergleich ≤, gültig ab 01.01.2025 laut Anmerkung"
        werte:
          steuern.ustg19.grenze_vorjahr: { version: 2 }                     # Zeile derselben ID, korrigierte Version
          gesellschaft.gmbh.mindesteinzahlung_gesamt: { version: 1, zeile: gesellschaft.gmbh.mindesteinzahlung_anteil }
  Ist die Version gleich der Version in der Liste, muss der Wert dem Katalog entsprechen (bestätigt wie vorgelegt, mit
  Anmerkung zu Vergleich oder Gültigkeit).
"""
import json
import os
import re
import subprocess
import sys

import openpyxl
import yaml

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import regelwerte as rw  # noqa: E402

ARG = sys.argv[1:]


def opt(name, standard=None):
    return ARG[ARG.index(name) + 1] if name in ARG else standard


if not ARG or ARG[0].startswith("--"):
    raise SystemExit(__doc__)
mappe, art, datum = ARG[0], opt("--art"), opt("--datum")
assert art in rw.ZAEHLT, "--art mail, blatt oder bericht"
assert datum and re.fullmatch(r"\d{4}-\d{2}-\d{2}", datum), "--datum JJJJ-MM-TT"
mit_anmerkung = {x.strip() for x in (opt("--mit-anmerkung", "") or "").split(",") if x.strip()}
PROBE = "--probe" in ARG
KORREKTUREN = opt("--korrekturen")

wb = openpyxl.load_workbook(mappe)
titel = str(wb["Anleitung"]["A1"].value or "")
rolle = next((r for r, name in rw.ROLLEN.items() if titel.endswith(name)), None)
assert rolle, f"Rolle aus dem Titel „{titel}“ nicht erkannt"
ws = wb["Regelwerte"]
kopf = [str(c.value or "") for c in ws[1]]
spalte = {k: kopf.index(k) for k in ("ID", "Version", "Wert + Einheit", "Anmerkung")}
spalte["Ergebnis"] = next(i for i, k in enumerate(kopf) if k.startswith("Ergebnis"))
katalog = {(d["id"], d["version"]): d for _, d in rw.laden()}

def wert_gleich(d, wert):
    """Wert + Einheit wie in der Liste (E-213: Jahreszahlen früher mit Tausenderpunkt, z. B. „2.027“)."""
    w = rw.wert_text(d["wert"], d.get("einheit", ""))
    return w == wert or (d.get("einheit") == "kalenderjahr" and w == wert.replace(".", ""))


zeilen = {}  # (ID, Version in der Liste) -> (Wert, Ergebnis, Anmerkung)
for z in ws.iter_rows(min_row=2, values_only=True):
    if z[spalte["ID"]]:
        zeilen[(z[spalte["ID"]], z[spalte["Version"]])] = (str(z[spalte["Wert + Einheit"]] or ""), str(z[spalte["Ergebnis"]] or "").strip(), str(z[spalte["Anmerkung"]] or "").strip())


def zeile_finden(ziel, f):
    """Zeile der Liste zu einem Eintrag in --korrekturen: `zeile: id` oder `zeile: id@version`; ohne Angabe dieselbe ID
    (gleiche Version, sonst die einzige Zeile dieser ID)."""
    z = str(f.get("zeile", ziel))
    if "@" in z:
        rid, v = z.split("@")
        k = (rid, int(v))
    elif (z, f.get("version")) in zeilen:
        k = (z, f.get("version"))
    else:
        treffer = [x for x in zeilen if x[0] == z]
        assert len(treffer) == 1, f"--korrekturen: Zeile zu {ziel} nicht eindeutig – zeile: <id>@<version> angeben"
        k = treffer[0]
    assert k in zeilen, f"--korrekturen: {k[0]}@{k[1]} steht nicht in der Liste"
    return k


richtig, falsch, offen, abweichend, anmerkung = [], [], [], [], []
spec = yaml.safe_load(open(KORREKTUREN, encoding="utf-8")) if KORREKTUREN else []
in_spec = {zeile_finden(ziel, f or {}) for g in spec for ziel, f in (g.get("werte") or {}).items()}
for z in ws.iter_rows(min_row=2, values_only=True):
    rid = z[spalte["ID"]]
    if not rid:
        continue
    ver, wert = z[spalte["Version"]], str(z[spalte["Wert + Einheit"]] or "")
    erg, anm = str(z[spalte["Ergebnis"]] or "").strip(), str(z[spalte["Anmerkung"]] or "").strip()
    d = katalog.get((rid, ver))
    if (rid, ver) in in_spec:
        continue
    if not erg:
        offen.append(rid)
    elif erg.lower().startswith("falsch"):
        falsch.append((rid, ver, wert, erg, anm))
    elif not erg.lower().startswith("richtig"):
        offen.append(f"{rid} (Ergebnis „{erg}“ unklar)")
    elif not d or not wert_gleich(d, wert):
        abweichend.append(f"{rid}@{ver}")
    elif rw.rolle(rid) != rolle:
        abweichend.append(f"{rid}@{ver} (gehört zur Rolle {rw.rolle(rid)})")
    elif anm and rid not in mit_anmerkung:
        anmerkung.append((rid, anm))
    else:
        richtig.append((rid, ver))

print(f"{rw.ROLLEN[rolle]}: {len(richtig)} richtig, {len(falsch)} falsch, {len(anmerkung)} mit Anmerkung, {len(abweichend)} abweichend, {len(offen)} ohne Ergebnis")
for rid, ver, wert, erg, anm in falsch:
    print(f"  falsch: {rid}@{ver} ({wert}) – {erg}{' – ' + anm if anm else ''}  → Wert als neue Version anlegen und erneut vorlegen")
for rid, anm in anmerkung:
    print(f"  Anmerkung lesen: {rid} – {anm}  → ohne Änderung übernehmen mit --mit-anmerkung {rid}")
for x in abweichend:
    print(f"  nicht übernommen (Wert, Version oder Rolle passt nicht mehr zum Katalog): {x}")
# E-213: Korrekturen bzw. präzisierte Zeilen laut --korrekturen (je Gruppe ein Block mit eigenem Umfang)
korrigiert = []  # (umfang, [(ziel, version, zeile, listen_version)])
for g in spec:
    umf = str(g.get("umfang") or "").strip()
    assert umf and "\n" not in umf, "--korrekturen: jede Gruppe braucht einen einzeiligen umfang"
    werte = []
    for ziel, f in (g.get("werte") or {}).items():
        f = f or {}
        version = f.get("version")
        zeile, lver = zeile_finden(ziel, f)
        lwert, erg, _anm = zeilen[(zeile, lver)]
        assert erg, f"--korrekturen: {zeile} ohne Ergebnis"
        assert (ziel, version) in katalog, f"--korrekturen: {ziel}@{version} gibt es nicht im Katalog (korrigierte Fassung zuerst anlegen)"
        assert rw.rolle(ziel) == rolle, f"--korrekturen: {ziel} gehört zur Rolle {rw.rolle(ziel)}"
        if ziel == zeile and version == lver:
            assert erg.lower().startswith("richtig") and wert_gleich(katalog[(ziel, version)], lwert), f"--korrekturen: {ziel}@{version} wie vorgelegt nur mit „richtig“ und gleichem Wert"
        else:
            assert version != lver or ziel != zeile, ziel
        werte.append((ziel, version, zeile, lver))
    korrigiert.append((umf, werte))
for umf, werte in korrigiert:
    print(f"  laut Anmerkung bzw. korrigiert ({umf}): " + ", ".join(f"{z}@{v}" + (f" (Liste: {zl}@{lv})" if (z, v) != (zl, lv) else "") for z, v, zl, lv in werte))
if not richtig and not korrigiert:
    raise SystemExit("Nichts zu übernehmen.")

# Register: neuer Block (mehrere Versionen derselben ID → je Version ein Block)
umfang = opt("--umfang", f"Regelwerte laut Prüfliste mit Ergebnis „richtig“ ({len(richtig)} Werte)")
assert "\n" not in umfang
gruppen = []
for rid, ver in richtig:
    for g in gruppen:
        if rid not in g:
            g[rid] = ver
            break
    else:
        gruppen.append({rid: ver})
quelle = "Bestätigung per Mail zur ausgefüllten Prüfliste Regelwerte (beim Gründer)" if art == "mail" else "ausgefüllte Prüfliste Regelwerte (beim Gründer)"
text = ""
quelle_k = quelle + "; korrigierte bzw. präzisierte Fassung laut Anmerkung des Prüfers"
for umf, werte in korrigiert:
    teile = []
    for z, v, zl, lv in werte:
        for g in teile:
            if z not in g:
                g[z] = (v, zl, lv)
                break
        else:
            teile.append({z: (v, zl, lv)})
    for g in teile:
        text += (f"- pruefer: {rolle}\n  umfang: alles\n  nachweis:\n    art: {art}\n    datum: '{datum}'\n    rolle: {rolle}\n"
                 f"    umfang: {json.dumps(umf, ensure_ascii=False)}\n    aufbewahrung: beim Gründer\n  eintraege:\n")
        for z, (v, zl, lv) in g.items():
            q = quelle if (z, v) == (zl, lv) else f"{quelle_k} (Liste: {zl} v{lv})"
            text += f"    {z}:\n      version: {v}\n      stand: '{datum}'\n      quelle: {json.dumps(q, ensure_ascii=False)}\n"
for g in gruppen:
    text += (f"- pruefer: {rolle}\n  umfang: alles\n  nachweis:\n    art: {art}\n    datum: '{datum}'\n    rolle: {rolle}\n"
             f"    umfang: {json.dumps(umfang, ensure_ascii=False)}\n    aufbewahrung: beim Gründer\n  eintraege:\n")
    for rid, ver in g.items():
        text += f"    {rid}:\n      version: {ver}\n      stand: '{datum}'\n      quelle: {json.dumps(quelle, ensure_ascii=False)}\n"
if PROBE:
    print(text)
    raise SystemExit(0)
alt = open(rw.REGISTER, encoding="utf-8").read()
neu = re.sub(r"^\[\]\s*$", "", alt, flags=re.M).rstrip("\n") + "\n" + text
open(rw.REGISTER, "w", encoding="utf-8").write(neu)

import anwenden  # noqa: E402  (prüft das ganze Register und schreibt die pruefung-Zeilen der Regelwerte)
uebersicht, geaendert, fehler = rw.anwenden(True, anwenden.nachweis_pruefen)
if fehler:
    open(rw.REGISTER, "w", encoding="utf-8").write(alt)
    rw.anwenden(True, anwenden.nachweis_pruefen)
    raise SystemExit("Register fehlerhaft, nichts übernommen:\n  " + "\n  ".join(fehler))
print("Regelwerte:", uebersicht, "geändert:", ", ".join(geaendert))
try:
    subprocess.run(["node", os.path.join(rw.ROOT, "apps/mobile/werkzeuge/katalog-buendel.mjs")], check=True)
except (subprocess.CalledProcessError, FileNotFoundError):
    print("Bündel der App bitte neu erzeugen: pnpm build && node apps/mobile/werkzeuge/katalog-buendel.mjs")
print("Danach: python3 docs/freigabe-katalog/anwenden.py --pruefen, node apps/mobile/werkzeuge/store-check.mjs")
