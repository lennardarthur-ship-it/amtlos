"""
Fachfragen an den Steuerberater zum Zugriff ohne Abo (GoBD), Stand 25.10.2026, als Excel-Datei (bis E-199 im Versandordner
versand-2026-10-25-steuern-gobd; das Paket selbst liegt nicht mehr im Repository).

    python3 docs/freigabe-katalog/gobd-fachfragen.py <ziel-ordner>     # außerhalb des Repositorys (E-199)
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ziel import ziel_ausserhalb  # noqa: E402  (E-199: Ziel außerhalb des Repositorys)

ZIEL = ziel_ausserhalb(sys.argv[1] if len(sys.argv) > 1 else None, "Aufruf: python3 docs/freigabe-katalog/gobd-fachfragen.py <ziel-ordner>")
os.makedirs(ZIEL, exist_ok=True)
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
wb=Workbook(); ws=wb.active; ws.title="Fachfragen"
ws["A1"]="Betriebsklar – Fachfragen zum Zugriff ohne Abo (GoBD)"; ws["A1"].font=Font(bold=True,size=14)
ws["A2"]="An: Steuerberater · Stand 25.10.2026 · Bitte je Frage Spalte D wählen und bei Bedarf Spalte E ausfüllen."
hdr=["Nr.","Frage","So ist es in der App","Antwort","Anmerkung / Formulierungsvorschlag"]
ws.append([]); ws.append(hdr)
for c in ws[4]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E5A"); c.alignment=Alignment(wrap_text=True,vertical="top")
rows=[
("1","Reicht es GoBD-rechtlich, dass nach Ende des Abos alle vorhandenen Aufzeichnungen nur noch lesbar und exportierbar sind, aber keine neuen mehr erfasst werden können? Muss darüber hinaus etwas zugänglich bleiben?",
 "Ohne Abo bleibt immer möglich: alle Rechnungen, Belege, Kassentage und das Kassenbuch ansehen; ausgegebene Rechnungen erneut teilen und drucken; Rechnungen stornieren und das Storno teilen; erfasste Kassentage mit Begründung berichtigen (wird protokolliert); Kassenbuch erfasster Tage ansehen, teilen, drucken, festschreiben; Belege exportieren; Sicherung, Wiederherstellen, Löschen. Nicht mehr möglich: neue Rechnungen, neue Belege, neue Tagesabschlüsse, Kassenstürze, Sachentnahmen. Es wird nichts gelöscht oder ausgeblendet."),
("2","Müssen wir Nutzer bei Ende des Abos ausdrücklich darauf hinweisen, dass ihre Aufzeichnungs- und Aufbewahrungspflichten weiterlaufen (insbesondere Kasse)? Falls ja: Passt unser Entwurf (rechts), oder wie soll er lauten?",
 "Bisher kein Hinweis bei Abo-Ende. Entwurf: „Dein Pro-Abo ist beendet. Deine Daten bleiben vollständig erhalten und lassen sich jederzeit ansehen, teilen und exportieren. Deine Aufzeichnungs- und Aufbewahrungspflichten laufen weiter – neue Rechnungen, Belege und Kassenaufzeichnungen musst du ab jetzt anders erfassen, bis du Pro wieder aktivierst.“"),
("3","Genügt nach Ende des Abos der kostenlose Export für den Datenzugriff der Finanzverwaltung (§ 147 Abs. 6 AO), oder müssen die maschinell auswertbaren Listen kostenlos bleiben?",
 "Ohne Abo: Sicherungsdatei (eigenes Format, vollständig), PDFs der Rechnungen, Belege einzeln, Kassenbuch als PDF. Nur mit Pro: Paket für die Steuerberatung (DATEV-Buchungsstapel, CSV-Listen Rechnungen/Belege/Kassenabschlüsse, Änderungsprotokoll). Unser Vorschlag: CSV-Listen und Änderungsprotokoll künftig auch ohne Abo, DATEV-Buchungsstapel bleibt Pro."),
]
for r in rows: ws.append(list(r)+["",""])
dv=DataValidation(type="list",formula1='"Ja, passt,Nein – siehe Anmerkung,Anpassen – siehe Anmerkung"',allow_blank=True); ws.add_data_validation(dv); dv.add("D5:D7")
thin=Side(style="thin",color="999999")
for row in ws.iter_rows(min_row=4,max_row=7,max_col=5):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
for col,w in zip("ABCDE",[5,45,60,22,45]): ws.column_dimensions[col].width=w
for r in (5,6,7): ws.row_dimensions[r].height=150
# E-197: keine Unterschrift – Bestätigung per E-Mail
ws["A10"]="Bitte bestätigen Sie Ihre Antworten per E-Mail (Antwort auf unsere Mail, mit dieser Datei im Anhang); eine Unterschrift ist nicht nötig."; ws["A10"].font=Font(bold=True)
ws.print_options.gridLines=False; ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToHeight=0
wb.save(os.path.join(ZIEL, "Betriebsklar-Fachfragen-GoBD-Zugriff-ohne-Abo-Steuerberater.xlsx"))
