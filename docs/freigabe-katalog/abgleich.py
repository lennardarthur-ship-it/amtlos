"""
Abgleich-Datei je Prüfer (E-173): alle Einträge seines Bereichs, die in der aktuellen Fassung noch nicht von ihm
freigegeben sind – mit Gegenüberstellung „Ihre geprüfte Fassung“ / „aktuelle Fassung“ für jedes geänderte Feld.

    python3 docs/freigabe-katalog/abgleich.py <ziel-ordner> [--notizen notizen.json]

notizen.json: {"<id>": "Satz, was sich geändert hat und warum"} (optional, erscheint je Eintrag).
"""
import glob
import json
import os
import sys

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HIER = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HIER, "../.."))
sys.argv_ = sys.argv
src = open(os.path.join(HIER, "anwenden.py"), encoding="utf-8").read().split("abweichung, uebersicht")[0]
g = {"__file__": os.path.join(HIER, "anwenden.py")}
exec(compile(src, "anwenden.py", "exec"), g)
gebiet, BENOETIGT, frei, fassung, gueltig, abgedeckt = g["gebiet"], g["BENOETIGT"], g["frei"], g["fassung"], g["gueltig"], g["abgedeckt"]
WEBER, BUSCH, KOCH, BAUM = g["WEBER"], g["BUSCH"], g["KOCH"], g["BAUM"]

ziel = os.path.abspath([a for a in sys.argv[1:] if not a.startswith("--")][0])
os.makedirs(ziel, exist_ok=True)
notizen = json.load(open(sys.argv[sys.argv.index("--notizen") + 1], encoding="utf-8")) if "--notizen" in sys.argv else {}
AUSGENOMMEN = set()  # E.56 seit E-174 bei RA Dr. Weber

FELDER = [("titel", "Titel"), ("kurz", "Kurztext"), ("erklaerung", "Erklärung"), ("ausloeser", "Auslöser"), ("frist", "Frist"),
          ("wiederholung", "Wiederholung"), ("rechtsgrundlage", "Rechtsgrundlage"), ("sanktion", "Folgen bei Verstoß"),
          ("stelle", "Zuständige Stelle"), ("status_recht", "Rechtsstand"), ("gueltig_ab", "Gültig ab"), ("quellen", "Quellen"),
          ("einordnung", "Einordnung"), ("erledigung", "Erledigung")]


def de(x):
    if isinstance(x, dict):
        return x.get("de") or ""
    return "" if x is None else str(x)


def ausloeser(a):
    if not a:
        return ""
    lab = f" „{de(a['label'])}“" if "label" in a else ""
    if a.get("immer"):
        return "gilt für alle Betriebe"
    for art, wort in (("all", " UND "), ("any", " ODER ")):
        if art in a:
            return "(" + wort.join(ausloeser(t) for t in a[art]) + ")" + lab
    if "not" in a:
        return "NICHT " + ausloeser(a["not"]) + lab
    if "feld" in a:
        w = a.get("wert")
        if isinstance(w, dict) and "regelwert" in w:
            w = "Regelwert " + w["regelwert"]
        return f"[{a['feld']} {a.get('op', '')} {'' if w is None else w}]".replace("  ", " ") + lab
    return str(a)


def frist(f):
    if not f:
        return ""
    t = f.get("typ")
    if t == "dauerhaft":
        return "laufend (keine feste Frist)"
    teile = [t]
    for k in ("bezug", "versatz", "zeitpunkt", "datum", "termin", "monat", "tag", "rrule"):
        if k in f:
            v = f[k]
            teile.append(f"{k}: {v.get('fakt', v) if isinstance(v, dict) else v}")
    return "; ".join(str(x) for x in teile)


def text(v, k=None):
    if v is None:
        return ""
    if k == "ausloeser":
        return ausloeser(v)
    if k == "frist":
        return frist(v)
    if isinstance(v, dict) and set(v) <= {"de", "en"}:
        return v.get("de", "")
    if isinstance(v, dict) and "text" in v:
        return f"{text(v['text'])} (Schwere: {v.get('schwere', '')})"
    if isinstance(v, list) and v and isinstance(v[0], dict) and "url" in v[0]:
        return "\n".join(f"{q.get('titel', '')} – {q['url']}" for q in v)
    if isinstance(v, list):
        return "; ".join(text(x) for x in v)
    return yaml.safe_dump(v, allow_unicode=True, default_flow_style=True, width=200).strip()


PERSON = {WEBER: "RA Dr. Ulrich Weber (Weber & Merz)", BUSCH: "StB Lutz Busch (Heckler und Busch)", KOCH: "Thomas Koch, Fachkraft für Arbeitssicherheit", BAUM: "Tim Baum, Lebensmittelfachbeauftragter"}
NAME = {WEBER: "Dr-Weber", BUSCH: "StB-Busch", KOCH: "Thomas-Koch", BAUM: "Tim-Baum"}
register = yaml.safe_load(open(os.path.join(ROOT, "docs/freigabe-katalog/freigaben.yaml"), encoding="utf-8"))
GESEHEN = {(r["pruefer"], r["umfang"]): r.get("gesehen") or {} for r in register}
je = {p: [] for p in PERSON}
for pfad in sorted(glob.glob(os.path.join(ROOT, "catalog/pflichten/*.yaml"))):
    d = yaml.safe_load(open(pfad, encoding="utf-8"))
    if d["id"] in AUSGENOMMEN:
        continue
    for p, u in BENOETIGT[gebiet(d)]:
        f = frei.get((p, u), {}).get(d["id"])
        if abgedeckt(pfad, d, p, u):
            continue
        if gueltig(pfad, d, f, u):
            # alles andere ist freigegeben – nur Erklärung und Stelle standen nicht in der Freigabe-Datei
            zeilen = [(lab, "(stand nicht in Ihrer Datei)", text(d.get(k), k)) for k, lab in FELDER if k in ("erklaerung", "stelle") and d.get(k)]
            je[p].append((d, f, zeilen, u, f"Inhalt von Ihnen freigegeben (v{f['version']}). Bitte nur Erklärung und Zuständige Stelle bestätigen."))
            continue
        gv = f["version"] if f else GESEHEN.get((p, u), {}).get(d["id"])
        alt = fassung(pfad, gv) if gv else None
        zeilen = []
        for k, label in FELDER:
            if u == "sanktion" and k != "sanktion":
                continue
            if u == "ohne_sanktion" and k == "sanktion":
                continue
            neu_t = text(d.get(k), k)
            alt_t = text(alt.get(k), k) if alt else None
            if alt is None or alt_t != neu_t:
                zeilen.append((label, alt_t if alt is not None else "(neu für Sie)", neu_t))
        if not gv:
            art = "Neu für Sie – bitte ganz prüfen."
        elif gv == d["version"]:
            art = "Gesehen, aber noch nicht freigegeben – bitte entscheiden."
        else:
            art = f"Geändert seit Ihrer Fassung v{gv} (meist Ihre Anmerkungen). Geänderte Felder unten."
        je[p].append((d, f or ({"version": gv} if gv else None), zeilen, u, art))

import re as _re
MARKE = json.load(open(os.path.join(ROOT, "packages/marke/marke.json"), encoding="utf-8"))
def marke(t):
    """wie im App-Bündel: Katalogtexte sagen „Amtlos“, angezeigt wird der Markenname"""
    return _re.sub(r"\bAmtlos\b", MARKE["name"], _re.sub(r"\bAmtlos Pro\b", MARKE["proName"], t)) if isinstance(t, str) else t


GELB = PatternFill("solid", fgColor="FFF4CC")
GRAU = PatternFill("solid", fgColor="EDEDED")
for p, liste in je.items():
    if not liste:
        continue
    wb = Workbook()
    ws = wb.active
    ws.title = "Abgleich"
    kopf = ["Nr", "ID", "Ihre Fassung", "Aktuell", "Was sich geändert hat (unsere Notiz)", "Feld", "In Ihrer Fassung", "Aktuelle Fassung", "Entscheidung", "Anmerkung"]
    ws.append(kopf)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = GRAU; c.alignment = Alignment(wrap_text=True, vertical="top")
    for d, f, zeilen, u, art in liste:
        erste = True
        for label, a, n in zeilen or [("–", "", "")]:
            ws.append([d.get("katalog_nr") if erste else "", d["id"] if erste else "", (f"v{f['version']}" if f else "–") if erste else "",
                       f"v{d['version']}" if erste else "", (art + (" " + notizen[d["id"]] if d["id"] in notizen else "")) if erste else "", label, marke(a), marke(n), "" , ""])
            r = ws.max_row
            if erste:
                ws.cell(r, 9).fill = GELB; ws.cell(r, 10).fill = GELB
            erste = False
    for i, b in enumerate([7, 32, 9, 8, 40, 16, 60, 60, 14, 40], start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = b
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"freigegeben,ändern"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"I2:I{ws.max_row}")
    ws.freeze_panes = "C2"
    hw = wb.create_sheet("So geht's")
    for z in [[f"Abgleich – {PERSON[p]}"], [""],
              [f"Diese Datei enthält die {len(liste)} Einträge Ihres Bereichs, die Sie in der aktuellen Fassung noch nicht freigegeben haben – weil sie sich nach Ihrer Durchsicht geändert haben (meist durch die Umsetzung Ihrer Anmerkungen) oder neu sind."],
              ["Je geändertem Feld stehen „In Ihrer Fassung“ und „Aktuelle Fassung“ nebeneinander. Unveränderte Felder sind nicht aufgeführt."],
              ["Bitte je Eintrag in der gelben Spalte „Entscheidung“ wählen: freigegeben (in der aktuellen Fassung) oder ändern."],
              ["Die Bestätigungsseite listet danach genau die Fassungen (Versionen), die Sie freigegeben haben."]]:
        hw.append(z)
    hw.column_dimensions["A"].width = 130
    hw["A1"].font = Font(bold=True, size=13)
    for r in hw.iter_rows():
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    datei = os.path.join(ziel, f"Betriebsklar-Abgleich-{NAME[p]}.xlsx")
    wb.save(datei)
    print(f"{PERSON[p]}: {len(liste)} Einträge → {datei}")
