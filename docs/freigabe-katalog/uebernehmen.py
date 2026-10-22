"""
Freigaben aus einer zurückgeschickten Freigabe-Mappe in den Katalog übernehmen (E-171).

    python3 docs/freigabe-katalog/uebernehmen.py <mappe.xlsx> "<Prüfer>" <JJJJ-MM-TT> [--umgesetzt id1,id2,...]

- Einträge mit Entscheidung „freigegeben“ werden freigegeben.
- Einträge mit „ändern“ nur, wenn ihre ID in --umgesetzt steht (Änderung umgesetzt und geprüft).
- „streichen“ und leere Zeilen bleiben unverändert.
Gesetzt wird `pruefung: {status: freigegeben, geprueft_von: [<Prüfer>], geprueft_am, geprueft_version: <aktuelle version>}`.
Danach: `node apps/mobile/werkzeuge/katalog-buendel.mjs` und `pnpm katalog:pruefen`.
"""
import os
import re
import sys

import openpyxl
import yaml

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "../.."))
mappe, pruefer, datum = sys.argv[1], sys.argv[2], sys.argv[3]
assert re.fullmatch(r"\d{4}-\d{2}-\d{2}", datum), "Datum als JJJJ-MM-TT"
umgesetzt = set()
if "--umgesetzt" in sys.argv:
    umgesetzt = {x.strip() for x in sys.argv[sys.argv.index("--umgesetzt") + 1].split(",") if x.strip()}

ws = openpyxl.load_workbook(mappe)["Einträge"]
kopf = [c.value for c in ws[1]]
i_id, i_ent = kopf.index("ID"), kopf.index("Entscheidung")
gesetzt, uebersprungen = [], []
for r in ws.iter_rows(min_row=2, values_only=True):
    pid, ent = r[i_id], (r[i_ent] or "").strip()
    if not pid:
        continue
    if not (ent == "freigegeben" or (ent == "ändern" and pid in umgesetzt)):
        uebersprungen.append((pid, ent or "leer"))
        continue
    pfad = os.path.join(ROOT, "catalog/pflichten", f"{pid}.yaml")
    text = open(pfad, encoding="utf-8").read()
    version = yaml.safe_load(text)["version"]
    block = f"pruefung:\n  status: freigegeben\n  geprueft_von:\n    - {pruefer}\n  geprueft_am: '{datum}'\n  geprueft_version: {version}\n"
    neu, n = re.subn(r"^pruefung:.*?(?=^\S|\Z)", block, text, count=1, flags=re.M | re.S)
    assert n == 1, f"kein pruefung-Block in {pid}"
    open(pfad, "w", encoding="utf-8").write(neu)
    gesetzt.append(pid)

print(f"freigegeben: {len(gesetzt)}")
for pid, ent in uebersprungen:
    print(f"  nicht freigegeben: {pid} ({ent})")
