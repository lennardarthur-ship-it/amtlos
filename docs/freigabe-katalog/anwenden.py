"""
Freigaben aus dem Register (docs/freigabe-katalog/freigaben.yaml) in den Katalog schreiben (E-173, E-196, E-197).

    python3 docs/freigabe-katalog/anwenden.py            # schreibt pruefung-Blöcke, druckt Übersicht
    python3 docs/freigabe-katalog/anwenden.py --pruefen  # schreibt nichts, Exit 1 wenn der Katalog vom Register abweicht
    python3 docs/freigabe-katalog/anwenden.py --bericht  # zusätzlich: je Eintrag, was ohne schriftliche Bestätigung freigegeben ist

Regel: Eine Freigabe gilt nur für die Fassung (Version), die der Prüfer gesehen hat. Ein Eintrag ist freigegeben, wenn
für seinen Bereich alle nötigen Freigaben zur aktuellen Version passen UND schriftlich bestätigt sind:
- Recht → `recht` (Rechtsanwalt); Steuern → `steuern` (Steuerberater); Arbeitsschutz → `arbeitsschutz` (Sifa);
- Lebensmittel → `hygiene` (Lebensmittelfachbeauftragter, Inhalt ohne Folgen) UND `recht` (nur Folgen bei Verstoß).
Schriftlich bestätigt (E-197, ersetzt die Unterschriftsregel aus E-196) heißt: Der Block hat einen `nachweis` mit `art`,
`datum` (Tag der Bestätigung), `rolle` (= `pruefer`) und `umfang` (was bestätigt ist, in unseren Worten). E-198: Mails der
Prüfer liegen nicht im Repository – nur der Gründer bewahrt sie auf (`aufbewahrung: beim Gründer`); kein Zitat aus einer Mail.
- `art: mail` – Bestätigung des Prüfers per E-Mail (Textform); mit `blatt` (die ausgefüllte Datei des Prüfers unter
  docs/pruefberichte) muss jeder Eintrag dort mit dieser Version „freigegeben“ sein;
- `art: unterschrift_entfernt` – das unterschriebene Blatt lag vor und wurde auf Wunsch des Gründers am 25.10.2026 aus dem
  Projekt entfernt (E-197); `umfang` sagt, was es bestätigte, `blatt` nennt die ausgefüllte Datei, soweit vorhanden;
- `art: blatt` – vom Prüfer ausgefüllte Freigabe-Datei (`blatt`, Spalte Entscheidung);
- `art: ersetzt` oder `offen` zählt nicht.
Das Datum im Vermerk (`geprueft_am`) liegt nie nach `datum`; bei `art: mail` ist es das Datum der Mail. Der pruefung-Block
trägt `bestaetigt_am`.
Bei Teil-Umfängen („ohne_sanktion“ / „sanktion“) bleibt eine Freigabe gültig, wenn sich seit der freigegebenen Version
nur Felder außerhalb ihres Umfangs geändert haben (Vergleich mit der Fassung aus der git-Historie).
Einschränkungen des Prüfers je Eintrag (`einschraenkung`, z. B. „ohne arbeitsmedizinische Angaben“) und offene
Prüfpunkte (docs/freigabe-katalog/offene-punkte.yaml) kommen in den pruefung-Block; die App zeigt sie im Vermerk.
Alles andere wird `in_pruefung` (zuvor oder ohne schriftliche Bestätigung freigegeben) bzw. bleibt `entwurf`. Die App zeigt
„geprüft“ nur bei `geprueft_version == version` (apps/mobile/src/pflichten/pruefer.ts).
"""
import glob
import json
import os
import re
import subprocess
import sys

import yaml

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "../.."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
PRUEFEN = "--pruefen" in sys.argv
BERICHT = "--bericht" in sys.argv

STEUER_GRUENDUNG = {"gruendung.fragebogen_steuerliche_erfassung", "gruendung.ust_idnr",
                    # Steuerberater 23.10.2026: Jahresabschluss der OHG, KG und e. K. (sein Vorschlag, er sieht ihn durch)
                    "gesellschaft.jahresabschluss_personenhandel"}
# E-195: elektronische Lohnsteuerbescheinigung (§ 41b EStG, § 93c AO) ist Steuerrecht – Steuerberater, obwohl Kategorie „personal“
STEUER_PERSONAL = {"personal.lohnsteuerbescheinigung"}
ARBEITSSCHUTZ_BRANCHE = {"branche.handwerk_gefaehrdung_gewerk", "branche.bau_baustellenverordnung", "branche.beauty_hautschutzplan", "branche.gastro_schankanlage_fluessiggas"}
LEBENSMITTEL_RECHT = {"branche.gastro_preisverzeichnis", "branche.gastro_mehrweg", "branche.gastro_ppwr_radar", "branche.gastro_einwegkunststofffonds"}
# Rollen statt Personen (E-196): Schlüssel in freigaben.yaml und in `pruefung.geprueft_von`
RECHT, STEUERN, SIFA, LMH = "recht", "steuern", "arbeitsschutz", "hygiene"
ROLLEN = {RECHT: "Rechtsanwalt", STEUERN: "Steuerberater", SIFA: "Fachkraft für Arbeitssicherheit", LMH: "Lebensmittelfachbeauftragter"}
NICHT_VERGLICHEN = {"version", "pruefung", "fachlich_zu_pruefen", "aenderungsnotiz"}
# Wörter in `fachlich_zu_pruefen`, die eine weitere fachliche Prüfung verlangen – jeder solche Punkt muss in
# offene-punkte.yaml als offen oder erledigt eingeordnet sein (E-196)
OFFEN_STICHWORTE = [r'bestätigen lassen', r'bitte ganz prüfen', r'bitte prüfen', r'Nummer und Höhe bestätigen', r'Betriebsarzt',
                    r'gegen das Amtsblatt', r'gegenlesen', r'nicht am Wortlaut gelesen', r'nicht gelesen\)', r'Tatbestände nicht gelesen',
                    r'nicht geprüft', r'ausstehend', r'\bOffen:', r'weiterhin offen', r'nötig\?', r'erneut vorlegen', r'vor de[mn] Launch',
                    r'nach Verkündung', r'nach Umsetzungsgesetz', r'fachlich festlegen', r'fachlich erstellen', r'fachlich prüfen',
                    r'fachlich zu bestätigen', r'nicht ergänzt', r'Platzhalter', r'nicht bestätigt', r'prüfen\)?$', r'am Wortlaut prüfen',
                    r'bitte .*prüfen', r'nicht wie angegeben übernommen', r'Höhe .*nicht g', r'erneut prüfen']
TECHNIK = [r'\(Technik', r'Technik\)']

# E-174: arbeitsmedizinische Vorsorge (ArbMedVV) prüft der Rechtsanwalt rechtlich – die Sifa hat sie als ärztliches Fachgebiet abgegeben
RECHT_STATT_ARBEITSSCHUTZ = {"arbeitsschutz.arbeitsmedizinische_vorsorge"}


def gebiet(d):
    i, k = d["id"], d["kategorie"]
    if i in RECHT_STATT_ARBEITSSCHUTZ:
        return "recht"
    if i in ARBEITSSCHUTZ_BRANCHE or k == "arbeitsschutz":
        return "arbeitsschutz"
    if k == "steuern" or i in STEUER_GRUENDUNG or i in STEUER_PERSONAL:
        return "steuer"
    if i.startswith("branche.gastro_") and i not in LEBENSMITTEL_RECHT:
        return "lebensmittel"
    return "recht"


BENOETIGT = {"recht": [(RECHT, "alles")], "steuer": [(STEUERN, "alles")], "arbeitsschutz": [(SIFA, "alles")],
             "lebensmittel": [(LMH, "ohne_sanktion"), (RECHT, "sanktion")]}

register = yaml.safe_load(open(os.path.join(ROOT, "docs/freigabe-katalog/freigaben.yaml"), encoding="utf-8"))


ZAEHLT = {"mail", "unterschrift_entfernt", "blatt"}
NACHWEIS_FELDER = {"art", "datum", "rolle", "umfang", "aufbewahrung", "blatt"}
# E-198: Mails der Prüfer (und unsere Mail-Entwürfe) gehören nicht ins Repository
MAILDATEI = re.compile(r"(\.eml|\.msg|\.mbox|mail[^/]*\.txt)$", re.I)
_blaetter = {}


def _blatt(pfad):
    """Einträge einer ausgefüllten Datei: ID → (Version, Entscheidung)."""
    if pfad not in _blaetter:
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(ROOT, pfad), data_only=True)
        out = {}
        for ws in wb.worksheets:
            zeilen = list(ws.iter_rows(values_only=True))
            if not zeilen:
                continue
            kopf = [str(c) if c else "" for c in zeilen[0]]
            if "ID" not in kopf or "Entscheidung" not in kopf or not ({"Aktuell", "Version"} & set(kopf)):
                continue
            i, e = kopf.index("ID"), kopf.index("Entscheidung")
            v = kopf.index("Aktuell") if "Aktuell" in kopf else kopf.index("Version")
            for z in zeilen[1:]:
                if z[i]:
                    out[str(z[i])] = (str(z[v]).lstrip("v"), str(z[e] or "").strip())
        _blaetter[pfad] = out
    return _blaetter[pfad]


def bestaetigt(r):
    """Tag der schriftlichen Bestätigung, wenn der Block einen gültigen Nachweis hat (E-197, E-198) – sonst None."""
    n = r.get("nachweis") or {}
    if n.get("art") not in ZAEHLT:
        return None
    if n.get("blatt") and not os.path.isfile(os.path.join(ROOT, n["blatt"])):
        raise SystemExit(f"freigaben.yaml: ausgefüllte Datei fehlt: {n['blatt']}")
    if n["art"] == "blatt" and not n.get("blatt"):
        raise SystemExit(f"freigaben.yaml: Block {r['pruefer']}/{r['umfang']} mit art: blatt ohne Datei")
    return str(n["datum"])


def nachweis_pruefen(r):
    """E-198: Jeder Block hat einen Nachweis mit art, datum, rolle und umfang – ohne Mail-Datei und ohne Zitat; jeder Eintrag
    eines bestätigten Blocks liegt nicht nach der Bestätigung und steht, wo eine ausgefüllte Datei genannt ist, dort mit
    dieser Version und „freigegeben“."""
    n = r.get("nachweis") or {}
    wo = f"freigaben.yaml: Block {r['pruefer']}/{r['umfang']}"
    for k in ("art", "datum", "rolle", "umfang"):
        assert n.get(k), f"{wo}: nachweis ohne {k}"
    assert set(n) <= NACHWEIS_FELDER, f"{wo}: unbekannte Felder im nachweis {sorted(set(n) - NACHWEIS_FELDER)} (keine Mail-Datei, kein Zitat – E-198)"
    assert n["rolle"] == r["pruefer"], f"{wo}: nachweis.rolle {n['rolle']} passt nicht"
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(n["datum"])), f"{wo}: nachweis.datum {n['datum']}"
    if n["art"] in ("mail", "unterschrift_entfernt"):
        assert n.get("aufbewahrung"), f"{wo}: Mail bzw. entferntes Blatt ohne Angabe, wo es aufbewahrt wird"
    if n.get("blatt"):
        assert not MAILDATEI.search(n["blatt"]), f"{wo}: {n['blatt']} ist eine Mail (E-198)"
    am = bestaetigt(r)
    if not am:
        return
    for i, f in r["eintraege"].items():
        assert "zitat" not in f, f"freigaben.yaml: {i} – kein Zitat aus einer Mail (E-198)"
        assert not MAILDATEI.search(str(f.get("quelle", ""))), f"freigaben.yaml: {i} – Quelle ist eine Mail-Datei (E-198)"
        assert str(f["stand"]) <= am, f"freigaben.yaml: {i} Stand {f['stand']} nach der Bestätigung {am}"
        if n["art"] == "mail":
            assert str(f["stand"]) == am, f"freigaben.yaml: {i} Stand {f['stand']} ist nicht das Datum der Mail {am}"
        if n.get("blatt"):
            v, e = _blatt(n["blatt"]).get(i, (None, None))
            assert v == str(f["version"]) and e == "freigegeben", f"freigaben.yaml: {i} v{f['version']} in {n['blatt']} nicht freigegeben ({v}, {e})"


def keine_mails():
    """E-198: keine Mail-Dateien im Repository (Mails der Prüfer, Entwürfe unserer Mails)."""
    dateien = subprocess.run(["git", "ls-files"], cwd=ROOT, capture_output=True, text=True).stdout.splitlines()
    treffer = [d for d in dateien if "node_modules" not in d and (d.startswith("docs/") and MAILDATEI.search(d) or re.search(r"(^|/)E-Mail[^/]*\.txt$", d))]
    for d in dateien:
        if d.startswith("docs/pruefberichte/") and d.endswith((".txt", ".md")) and os.path.isfile(os.path.join(ROOT, d)):
            kopf = open(os.path.join(ROOT, d), encoding="utf-8", errors="replace").read(2000)
            if re.search(r"^(Betreff|Subject):", kopf, re.M) and re.search(r"^(Von|From|An|To|Anh(ä|ae)nge?):", kopf, re.M):
                treffer.append(d)
    return treffer


# mehrere Blöcke je Prüfer/Umfang möglich (z. B. Nachtrag, der Erklärung und Stelle bestätigt)
BLOECKE = [(r["pruefer"], r["umfang"], set(r.get("felder_fehlen") or []), r["eintraege"], bestaetigt(r)) for r in register]
for _r in register:
    assert _r["pruefer"] in ROLLEN, f"freigaben.yaml: unbekannte Rolle {_r['pruefer']}"
    assert "unterschrift" not in _r and "unterschrieben_am" not in _r, "freigaben.yaml: Unterschriftsfelder durch `nachweis` ersetzt (E-197)"
    nachweis_pruefen(_r)
frei = {}
for _p, _u, _fehl, _e, _am in BLOECKE:
    for _i, _f in _e.items():
        frei.setdefault((_p, _u), {})[_i] = _f  # jüngster Block gewinnt (für Anzeige/Abgleich)
SONDERFELDER = {"erklaerung", "stelle"}

_offen_roh = yaml.safe_load(open(os.path.join(ROOT, "docs/freigabe-katalog/offene-punkte.yaml"), encoding="utf-8"))
OFFEN = _offen_roh.get("offen") or []
ERLEDIGT = _offen_roh.get("erledigt") or []

_hist = {}


def fassung(pfad, version):
    """Inhalt des Eintrags in der genannten Version aus der git-Historie (jüngster Stand mit dieser Version)."""
    if pfad not in _hist:
        rel = os.path.relpath(pfad, ROOT)
        commits = subprocess.run(["git", "log", "--format=%H", "--", rel], cwd=ROOT, capture_output=True, text=True).stdout.split()
        _hist[pfad] = []
        for c in commits:
            t = subprocess.run(["git", "show", f"{c}:{rel}"], cwd=ROOT, capture_output=True, text=True).stdout
            try:
                _hist[pfad].append(yaml.safe_load(t))
            except yaml.YAMLError:
                pass
    for d in _hist[pfad]:
        if d and d.get("version") == version:
            return d
    return None


def gueltig(pfad, d, f, umfang):
    if not f:
        return False
    if f["version"] == d["version"]:
        return True
    if umfang == "alles":
        return False
    alt = fassung(pfad, f["version"])
    if not alt:
        return False
    felder = (set(alt) | set(d)) - NICHT_VERGLICHEN
    if umfang == "ohne_sanktion":
        felder -= {"sanktion"}
    elif umfang == "sanktion":
        felder = {"sanktion"}
    return all(alt.get(k) == d.get(k) for k in felder)


def deckung(pfad, d, p, u, nur_bestaetigt):
    """(abgedeckt, Stand, Tag der Bestätigung, Einschränkungen) für Prüfer p mit Umfang u in der aktuellen Version."""
    noetig = {"rest"} if u == "sanktion" else {"rest"} | SONDERFELDER
    gedeckt, staende, bestaetigungen, einschr = set(), [], [], []
    for bp, bu, fehl, e, am in BLOECKE:
        if (bp, bu) != (p, u) or (nur_bestaetigt and not am):
            continue
        f = e.get(d["id"])
        if gueltig(pfad, d, f, u):
            gedeckt |= ({"rest"} | SONDERFELDER) - fehl
            staende.append(str(f["stand"]))
            if am:
                bestaetigungen.append(am)
            if f.get("einschraenkung"):
                einschr.append(f["einschraenkung"])
    ok = noetig <= gedeckt
    return ok, (max(staende) if ok else None), (max(bestaetigungen) if ok and bestaetigungen else None), sorted(set(einschr))


def abgedeckt(pfad, d, p, u, nur_bestaetigt=False):
    """Sind alle Felder im Umfang für die aktuelle Version freigegeben (inkl. Erklärung und Stelle)?"""
    return deckung(pfad, d, p, u, nur_bestaetigt)[0]


def offene_punkte(d):
    """Themen offener Prüfpunkte des Eintrags; prüft, dass jeder Stichwort-Punkt eingeordnet ist."""
    themen, fehlend = [], []
    for x in d.get("fachlich_zu_pruefen") or []:
        x = str(x)
        if any(re.search(t, x) for t in TECHNIK) or not any(re.search(o, x) for o in OFFEN_STICHWORTE):
            continue
        anfang = x[:60]
        o = [p for p in OFFEN if p["id"] == d["id"] and p["punkt"] == anfang]
        if o:
            themen.append(o[0]["thema"])
        elif not any(p["id"] == d["id"] and p["punkt"] == anfang for p in ERLEDIGT):
            fehlend.append(anfang)
    return themen, fehlend


def auswerten(pfad, d):
    """Prüfstand eines Eintrags: Status, Block für den Katalog und Angaben für Bericht und Nachprüfliste."""
    noetig = BENOETIGT[gebiet(d)]
    signiert = [(p, u) + deckung(pfad, d, p, u, True)[1:] + (deckung(pfad, d, p, u, True)[0],) for p, u in noetig]
    themen, fehlend = offene_punkte(d)
    info = {"id": d["id"], "version": d["version"], "rollen": [p for p, _ in noetig], "offen": themen, "nicht_eingeordnet": fehlend,
            "ohne_nachweis": [], "einschraenkung": []}
    for p, u in noetig:
        if not abgedeckt(pfad, d, p, u, True) and abgedeckt(pfad, d, p, u, False):
            info["ohne_nachweis"].append({"rolle": p, "umfang": u})
    if all(s[-1] for s in signiert):
        # geteilte Prüfung: Stand nie nach der Bestätigung einer der beteiligten Rollen (E-197)
        stand = min(s[2] for s in signiert)
        am = max(s[3] for s in signiert)
        einschr = sorted({e for s in signiert for e in s[4]})
        info["einschraenkung"] = einschr
        block = ("pruefung:\n  status: freigegeben\n  geprueft_von:\n" + "".join(f"    - {p}\n" for p, _ in noetig)
                 + f"  geprueft_am: '{stand}'\n  geprueft_version: {d['version']}\n  bestaetigt_am: '{am}'\n")
        if einschr:
            block += f"  einschraenkung: {json.dumps('; '.join(einschr), ensure_ascii=False)}\n"
        if themen:
            block += "  offen:\n" + "".join(f"    - {json.dumps(t, ensure_ascii=False)}\n" for t in themen)
        status = "freigegeben"
    else:
        irgendeine = any(frei.get((p, u), {}).get(d["id"]) for p, u in noetig)
        status = "in_pruefung" if irgendeine else "entwurf"
        block = f"pruefung:\n  status: {status}\n  geprueft_von: null\n  geprueft_am: null\n"
    info["status"] = status
    return status, block, info


def vorlagen_nachweis_pruefen():
    """E-197, E-198: Jede Vorlagen-Freigabe in geprueft.ts hat einen Nachweis mit art, datum (= bestaetigt_am), rolle (= Rolle der
    Freigabe) und umfang – ohne Mail-Datei und ohne Zitat; bei `mail` die Aufbewahrung, bei `bericht` den Bericht im
    Repository. Nennt die Freigabe eine ausgefüllte Datei (`blatt`), steht dort die bestätigte Prüfsumme (`pruefsumme_alt`,
    12 Zeichen) mit „freigegeben“ (bzw. im PDF des Blatts)."""
    import openpyxl
    ts = open(os.path.join(ROOT, "packages/fachlogik/src/dokumente/geprueft.ts"), encoding="utf-8").read()
    fehler = []
    for zeile in ts.splitlines():
        m = re.search(r'pruefsumme_alt: "([0-9a-f]{12})', zeile)
        if not m or "stand:" not in zeile:
            continue
        kurz = m.group(1)
        stand = re.search(r'stand: "([\d-]+)"', zeile).group(1)
        rolle = re.search(r'stand: "[\d-]+", rolle: "(\w+)"', zeile).group(1)
        am = re.search(r'bestaetigt_am: ("[\d-]+"|null)', zeile)
        n = re.search(r'nachweis: \{ art: "(\w+)", datum: "([\d-]+)", rolle: "(\w+)", umfang: "((?:[^"\\]|\\.)+)"(?:, aufbewahrung: "([^"]+)")?(?:, beleg: "([^"]+)")? \}', zeile)
        blatt = re.search(r'\bblatt: "([^"]+)"', zeile)
        blatt = blatt.group(1) if blatt else None
        if not am or am.group(1) == "null" or not n:
            fehler.append(f"{kurz}: ohne Nachweis (art, datum, rolle, umfang)")
            continue
        art, datum, nrolle, umfang, aufbewahrung, beleg = n.groups()
        am = am.group(1).strip('"')
        if datum != am:
            fehler.append(f"{kurz}: nachweis.datum {datum} ≠ bestaetigt_am {am}")
        if nrolle != rolle:
            fehler.append(f"{kurz}: nachweis.rolle {nrolle} ≠ {rolle}")
        if stand > am:
            fehler.append(f"{kurz}: Stand {stand} nach der Bestätigung {am}")
        if art == "mail" and not aufbewahrung:
            fehler.append(f"{kurz}: Mail ohne Angabe, wo sie aufbewahrt wird")
        if art == "bericht" and not (beleg and os.path.isfile(os.path.join(ROOT, beleg))):
            fehler.append(f"{kurz}: Bericht fehlt ({beleg})")
        if art not in ("mail", "bericht"):
            fehler.append(f"{kurz}: unbekannte Art {art}")
        if "nachweis_text" in zeile or MAILDATEI.search(beleg or "") or MAILDATEI.search(blatt or "") or re.search(r'docs/pruefberichte/[^"]*(mail[^"/]*\.txt|\.eml)', zeile):
            fehler.append(f"{kurz}: Mail-Datei oder Mail-Text in geprueft.ts (E-198)")
        if not blatt:
            continue
        if not os.path.isfile(os.path.join(ROOT, blatt)):
            fehler.append(f"{kurz}: {blatt} fehlt")
            continue
        if blatt.endswith(".pdf"):
            text = subprocess.run(["pdftotext", os.path.join(ROOT, blatt), "-"], capture_output=True, text=True).stdout
            if kurz not in text:
                fehler.append(f"{kurz}: nicht in {blatt}")
            continue
        wb = openpyxl.load_workbook(os.path.join(ROOT, blatt), data_only=True)
        ok = False
        for ws in wb.worksheets:
            zeilen = list(ws.iter_rows(values_only=True))
            kopf = [str(c) if c else "" for c in (zeilen[0] if zeilen else [])]
            if "Prüfsumme (Kurz)" in kopf and "Entscheidung" in kopf:
                ps, en = kopf.index("Prüfsumme (Kurz)"), kopf.index("Entscheidung")
                ok |= any(str(z[ps]) == kurz and str(z[en]).strip() == "freigegeben" for z in zeilen[1:])
        if not ok:
            fehler.append(f"{kurz}: in {blatt} nicht mit „freigegeben“")
    return fehler


def alle():
    for pfad in sorted(glob.glob(os.path.join(ROOT, "catalog/pflichten/*.yaml"))):
        text = open(pfad, encoding="utf-8").read()
        yield pfad, text, yaml.safe_load(text)


if __name__ == "__main__":
    abweichung, uebersicht, berichte, ungeordnet = [], {"freigegeben": 0, "in_pruefung": 0, "entwurf": 0}, [], []
    for pfad, text, d in alle():
        status, block, info = auswerten(pfad, d)
        uebersicht[status] += 1
        berichte.append(info)
        ungeordnet += [f"{d['id']}: {x}" for x in info["nicht_eingeordnet"]]
        neu, n = re.subn(r"^pruefung:.*?(?=^[^\s#-]|\Z)", block, text, count=1, flags=re.M | re.S)
        assert n == 1, d["id"]
        if neu != text:
            abweichung.append(d["id"])
            if not PRUEFEN:
                open(pfad, "w", encoding="utf-8").write(neu)

    je_rolle = {r: sum(1 for b in berichte if b["status"] == "freigegeben" and r in b["rollen"]) for r in ROLLEN}
    print(uebersicht, f"geändert: {len(abweichung)}", "je Rolle:", je_rolle)
    if BERICHT:
        for b in berichte:
            if b["ohne_nachweis"] or b["offen"] or b["einschraenkung"]:
                print(json.dumps(b, ensure_ascii=False))
    if ungeordnet:
        print("Prüfpunkte ohne Einordnung in offene-punkte.yaml:", *ungeordnet, sep="\n  ")
        sys.exit(1)
    mails = keine_mails()
    if mails:
        print("Mail-Dateien im Repository (E-198: Mails gehören nicht ins Projekt):", *mails, sep="\n  ")
        sys.exit(1)
    vfehler = vorlagen_nachweis_pruefen()
    if vfehler:
        print("Vorlagen-Freigaben ohne passenden Nachweis (geprueft.ts):", *vfehler, sep="\n  ")
        sys.exit(1)
    if PRUEFEN and abweichung:
        print("Katalog weicht vom Freigabe-Register ab:", ", ".join(abweichung[:20]))
        sys.exit(1)
