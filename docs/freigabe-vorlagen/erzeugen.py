"""
Freigabe der Vorlagen (T-52 ff.): je Prüfer ein Paket mit den Vorlagen-PDFs (leer, Auswahl, voll), einer Freigabe-Datei
(xlsx: je Vorlage freigegeben / ändern / nicht mein Bereich) und einem Bestätigungsblatt mit Prüfsumme je Vorlage.
Die Prüfsumme (vorlagenPruefsumme) hält die Fassung fest: Nach der Rückmeldung wird sie in
packages/fachlogik/src/dokumente/geprueft.ts eingetragen; ändert sich danach ein Text, verschwindet der Vermerk.

    python3 docs/freigabe-vorlagen/erzeugen.py <mappen-ordner> <ziel>
    # <mappen-ordner>: Ausgabe von docs/pruefmappe-{anwalt,arbeitsschutz,steuer}/erzeugen.mjs in anwalt/, arbeitsschutz/, steuer/
"""
import csv, hashlib, html, json, os, shutil, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "../.."))
MAPPEN, ZIEL = sys.argv[1], sys.argv[2]
# --nur '{"1 RA Dr. Weber": ["handelsregister", "Nutzungsbedingungen (Entwurf)"], ...}': nur diese Vorlagen/Texte (Nachprüfung)
NUR = json.loads(sys.argv[sys.argv.index("--nur") + 1]) if "--nur" in sys.argv else None
RUNDE = " (Nachprüfung)" if NUR else ""
NOTIZ = json.load(open(sys.argv[sys.argv.index("--notizen") + 1], encoding="utf-8")) if "--notizen" in sys.argv else {}
STAND = "23.10.2026"

ZUORDNUNG = {
    "1 RA Dr. Weber": {
        "name": "RA Dr. Ulrich Weber, Weber & Merz, Reutlingen",
        "vorlagen": ["datenschutz_dokumentation", "datenschutzerklaerung_website", "avv", "datenschutzinformation", "datenpanne", "video_hinweisschild",
                     "anamnese_einwilligung", "arbeitszeitnachweis", "nachweis_arbeitsbedingungen", "ki_kompetenz", "ausweis_hinweis", "schluesselquittung",
                     "impressum", "gewerbeanmeldung", "handelsregister", "transparenzregister"],
        "umfang": "rechtlich (Arbeitsrecht, Datenschutz, Gesellschafts- und Gewerberecht)",
        "texte": [("Datenschutzhinweis der App", "anwalt/datenschutzhinweis-app.md", "Datenschutzhinweis-App.pdf"),
                  ("Nutzungsbedingungen (Entwurf)", "anwalt/nutzungsbedingungen-entwurf.md", "Nutzungsbedingungen-Entwurf.pdf"),
                  ("Datenschutzerklärung der Website", "anwalt/website-datenschutz.html", "Website-Datenschutzerklaerung.pdf"),
                  ("Impressum der Website", "anwalt/website-impressum.html", "Website-Impressum.pdf")],
    },
    "2 StB Busch": {
        "name": "StB Lutz Busch, Kanzlei Heckler und Busch",
        "vorlagen": ["rechnung", "euer", "bilanz_guv", "ust_voranmeldung", "oss_meldung", "zm_meldung", "steuererklaerung", "verfahrensdokumentation",
                     "kassenbuch", "kassenmeldung", "aufbewahrung", "lohnkonto", "sv_meldedaten"],
        "umfang": "steuerlich (Umsatzsteuer, GoBD, Kasse, Lohnkonto und SV-Meldedaten)",
    },
    "3 Thomas Koch": {
        "name": "Thomas Koch, Fachkraft für Arbeitssicherheit",
        "vorlagen": ["gefaehrdungsbeurteilung", "unterweisung", "gefahrstoffverzeichnis", "betriebsanweisung", "betriebsanweisung_maschine",
                     "betriebsanweisung_biostoffe", "expositionsverzeichnis", "pruefprotokoll_elektro", "hautschutzplan", "hygieneplan"],
        "umfang": "fachlich im Arbeitsschutz",
        "hinweis": {"hygieneplan": "Nur die Fassung „Friseur, Kosmetik, Nagelstudio, Fußpflege“ – die Lebensmittel-Fassung ist schon geprüft. Infektionshygiene nach Landes-Hygieneverordnung: Wenn das nicht Ihr Fachgebiet ist, bitte „nicht mein Bereich“ wählen."},
    },
}

# Titel und Prüfsummen aus der Fachlogik (vorher: pnpm --filter @amtlos/fachlogik build)
js = """import(process.argv[1]).then(F=>{const o={};for(const v of F.VORLAGEN){o[v.id]={titel:v.titel,bereich:v.bereich,summe:F.vorlagenPruefsumme(v,v.id==='hygieneplan'?{fassung:{feld:'art',wert:'beauty'}}:undefined)}};console.log(JSON.stringify(o))})"""
V = json.loads(subprocess.run(["node", "-e", js, os.path.join(ROOT, "packages/fachlogik/dist/index.js")], capture_output=True, text=True, check=True).stdout)
json.dump({k: v["summe"] for k, v in V.items()}, open(os.path.join(ZIEL, "pruefsummen.json") if os.path.isdir(ZIEL) else "/dev/null", "w"))

def pdf(html_text, ziel):
    tmp = ziel[:-4] + ".html"
    open(tmp, "w", encoding="utf-8").write(html_text)
    s = f"""const {{ chromium }} = require('playwright-core');(async()=>{{const b=await chromium.launch({{executablePath:'/opt/pw-browsers/chromium',args:['--no-sandbox']}});const p=await b.newPage();await p.goto('file://{tmp}');await p.pdf({{path:'{ziel}',format:'A4',preferCSSPageSize:true,margin:{{top:'14mm',bottom:'12mm',left:'14mm',right:'14mm'}}}});await b.close();}})();"""
    subprocess.run(["node", "-e", s], check=True, cwd=os.path.join(ROOT, "apps/web"))
    os.remove(tmp)

CSS = "body{font-family:'Helvetica Neue',Arial,sans-serif;font-size:9.5pt;color:#111}h1{font-size:14pt;margin:0 0 2mm}p{margin:1.6mm 0;line-height:1.35}table{border-collapse:collapse;width:100%;font-size:8.5pt;margin-top:3mm}td,th{border:.4pt solid #999;padding:1.2mm 1.6mm;text-align:left;vertical-align:top}th{background:#eee}.sig{display:flex;gap:10mm;margin-top:12mm}.sig div{flex:1;border-top:.6pt solid #111;padding-top:1mm;font-size:8.5pt}.box{border:.4pt solid #888;padding:2.5mm 3mm;margin-top:3mm}"

def md_zu_pdf(md_pfad, ziel, titel):
    import re
    t = open(md_pfad, encoding="utf-8").read()
    try:
        import markdown
        body = markdown.markdown(t, extensions=["tables"])
    except ImportError:
        body = "<pre style='white-space:pre-wrap;font-family:inherit'>" + html.escape(t) + "</pre>"
    pdf(f"<!doctype html><meta charset=utf-8><title>{html.escape(titel)}</title><style>{CSS}</style>{body}", ziel)

def html_zu_pdf(pfad, ziel):
    tmp = ziel[:-4] + ".src.html"
    shutil.copy(pfad, tmp)
    s = f"""const {{ chromium }} = require('playwright-core');(async()=>{{const b=await chromium.launch({{executablePath:'/opt/pw-browsers/chromium',args:['--no-sandbox']}});const p=await b.newPage();await p.goto('file://{tmp}');await p.pdf({{path:'{ziel}',format:'A4',margin:{{top:'14mm',bottom:'12mm',left:'14mm',right:'14mm'}}}});await b.close();}})();"""
    subprocess.run(["node", "-e", s], check=True, cwd=os.path.join(ROOT, "apps/web"))
    os.remove(tmp)

if NUR:
    for o in list(ZUORDNUNG):
        if o not in NUR:
            del ZUORDNUNG[o]; continue
        ZUORDNUNG[o]["vorlagen"] = [v for v in ZUORDNUNG[o]["vorlagen"] if v in NUR[o]]
        ZUORDNUNG[o]["texte"] = [t for t in ZUORDNUNG[o].get("texte", []) if t[0] in NUR[o]]
GELB = PatternFill("solid", fgColor="FFF4CC"); GRAU = PatternFill("solid", fgColor="EDEDED")
for ordner, z in ZUORDNUNG.items():
    o = os.path.join(ZIEL, ordner); os.makedirs(os.path.join(o, "Vorlagen-PDF"), exist_ok=True)
    quelle = os.path.join(MAPPEN, "anwalt/vorlagen")
    for vid in z["vorlagen"]:
        for f in sorted(os.listdir(quelle)):
            if f.startswith(vid + "-") and f.endswith(".pdf") and not any(f.startswith(a + "-") for a in V if a != vid and a.startswith(vid + "_")):
                shutil.copy(os.path.join(quelle, f), os.path.join(o, "Vorlagen-PDF", f))
    wb = Workbook(); ws = wb.active; ws.title = "Freigabe Vorlagen"
    kopf = ["Nr", "Vorlage (ID)", "Titel", "PDFs im Ordner Vorlagen-PDF", "Prüfsumme (Kurz)", "Hinweis", "Entscheidung", "Anmerkung"]
    ws.append(kopf)
    for c in ws[1]: c.font = Font(bold=True); c.fill = GRAU
    for i, vid in enumerate(z["vorlagen"], 1):
        pdfs = ", ".join(sorted(f for f in os.listdir(os.path.join(o, "Vorlagen-PDF")) if f.startswith(vid + "-")))
        ws.append([i, vid, V[vid]["titel"], pdfs, V[vid]["summe"][:12], NOTIZ.get(vid) or z.get("hinweis", {}).get(vid, ""), "", ""])
        ws.cell(ws.max_row, 7).fill = GELB; ws.cell(ws.max_row, 8).fill = GELB
    texte = [(t, q, f, hashlib.sha256(open(os.path.join(MAPPEN, q), "rb").read()).hexdigest()) for t, q, f in z.get("texte", [])]
    for t, q, f, h in texte:
        ws.append([ws.max_row, "(Text)", t, "App- und Website-Texte/" + f, h[:12], (NOTIZ.get(t) or "Fassung aus dem Code vom 23.10.2026") + ("; Abschnitt Hosting folgt nach Wahl des Hosters" if "website-datenschutz" in q else ""), "", ""])
        ws.cell(ws.max_row, 7).fill = GELB; ws.cell(ws.max_row, 8).fill = GELB
    for i, b in enumerate([5, 28, 34, 44, 15, 40, 18, 50], 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = b
    for r in ws.iter_rows():
        for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"freigegeben,ändern,nicht mein Bereich"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"G2:G{ws.max_row}")
    ws.freeze_panes = "C2"
    if ordner.startswith("3") and os.path.exists(os.path.join(MAPPEN, "arbeitsschutz/bausteine.csv")):
        bs = wb.create_sheet("Bausteine (zum Kommentieren)")
        for zeile in csv.reader(open(os.path.join(MAPPEN, "arbeitsschutz/bausteine.csv"), encoding="utf-8-sig"), delimiter=";"): bs.append(zeile)
        for c in bs[1]: c.font = Font(bold=True); c.fill = GRAU
        bs.freeze_panes = "B2"
    hw = wb.create_sheet("So geht's")
    for t in [f"Freigabe der Vorlagen – {z['name']}", "",
              "Je Vorlage liegen im Ordner „Vorlagen-PDF“ drei Fassungen: leer (zum Ausfüllen), Auswahl (jede Auswahl umgestellt, damit alle Abschnitte erscheinen) und voll (realistisch ausgefüllt, Beispieldaten erfunden).",
              "Bitte je Vorlage in Spalte G wählen: freigegeben, ändern (Anmerkung in Spalte H) oder nicht mein Bereich.",
              "Die Prüfsumme hält genau diese Fassung fest. Nur solange sich an Texten und Bausteinen nichts ändert, zeigt die App bei der Vorlage „Vorlagen fachlich durchgesehen (Stand …)“ – ohne Namen. Jede Änderung hebt den Vermerk auf, bis Sie erneut durchgesehen haben.",
              "Die App füllt nichts für den Nutzer aus und berät nicht im Einzelfall; sie bietet Vorlagen und vorformulierte Vorschläge, die der Betrieb selbst übernimmt und anpasst."]:
        hw.append([t])
    hw.column_dimensions["A"].width = 130; hw["A1"].font = Font(bold=True, size=13)
    for r in hw.iter_rows():
        for c in r: c.alignment = Alignment(wrap_text=True)
    kurz = ordner.split(" ", 1)[1].replace(" ", "-").replace(".", "")
    wb.save(os.path.join(o, f"Betriebsklar-Freigabe-Vorlagen-{kurz}.xlsx"))
    zeilen = "".join(f"<tr><td>{html.escape(V[v]['titel'])}</td><td>{v}</td><td><code>{V[v]['summe'][:12]}</code></td></tr>" for v in z["vorlagen"])
    zeilen += "".join(f"<tr><td>{html.escape(t)}</td><td>Text</td><td><code>{h[:12]}</code></td></tr>" for t, q, f, h in texte)
    anzahl = len(z["vorlagen"]) + len(texte)
    blatt = f"""<!doctype html><meta charset=utf-8><style>@page{{size:A4;margin:14mm}}{CSS}</style>
<h1>Bestätigung der fachlichen Durchsicht – Vorlagen der App Betriebsklar{RUNDE}</h1>
<p>Anbieter: Lang Ventures UG (haftungsbeschränkt) i. G., Gomaringer Str. 33, 72770 Reutlingen</p>
<p><b>{html.escape(z['name'])}</b> – Stand der Durchsicht: ____________</p>
<p>Ich habe die unten aufgeführten {anzahl} Vorlagen{" und Texte" if texte else ""} in der Fassung der mir übergebenen PDFs (leer, Auswahl, voll) und der Freigabe-Datei vom {STAND} {html.escape(z['umfang'])} durchgesehen. Die Fassung ist durch die angegebene Prüfsumme bestimmt. Ausgenommen sind Vorlagen, die ich in der Freigabe-Datei mit „ändern“ oder „nicht mein Bereich“ markiert habe. Die Durchsicht betrifft die Vorlagen als allgemeine Muster, nicht ihre Verwendung im Einzelfall; spätere Rechtsänderungen sind nicht erfasst. Ändert sich eine Vorlage (andere Prüfsumme), gilt diese Bestätigung für sie nicht mehr, bis ich sie erneut durchgesehen habe.</p>
<table><tr><th>Vorlage</th><th>ID</th><th>Prüfsumme</th></tr>{zeilen}</table>
<div class="box"><b>Anzeige in der App:</b> „Vorlagen fachlich durchgesehen (Stand …)“ – ohne Namen, nur bei Vorlagen mit passender Prüfsumme.</div>
<div class="sig"><div>Ort, Datum</div><div>Unterschrift – {html.escape(z['name'])}</div></div>"""
    pdf(blatt, os.path.join(o, f"Betriebsklar-Bestaetigung-Vorlagen-{kurz}.pdf"))
    print(ordner, len(z["vorlagen"]), "Vorlagen,", len(os.listdir(os.path.join(o, "Vorlagen-PDF"))), "PDFs")

# Weber zusätzlich: Datenschutzhinweis der App, Nutzungsbedingungen, Website-Texte
A = os.path.join(MAPPEN, "anwalt")
if "1 RA Dr. Weber" in ZUORDNUNG and ZUORDNUNG["1 RA Dr. Weber"].get("texte"):
    w = os.path.join(ZIEL, "1 RA Dr. Weber", "App- und Website-Texte"); os.makedirs(w, exist_ok=True)
    for t, q, f in ZUORDNUNG["1 RA Dr. Weber"]["texte"]:
        (md_zu_pdf(os.path.join(MAPPEN, q), os.path.join(w, f), t) if q.endswith(".md") else html_zu_pdf(os.path.join(MAPPEN, q), os.path.join(w, f)))
    print("Weber: App- und Website-Texte", os.listdir(w))
