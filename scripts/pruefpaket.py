#!/usr/bin/env python3
"""
Prüfpaket für die fachliche Freigabe (E-143): alle Pflichten und Regelwerte des Katalogs als Excel-Tabelle für
Anwältin, Anwalt oder Steuerberatung – mit Auslösern in Klartext, Fristen, Rechtsgrundlagen, Quellen und den
offenen Prüfpunkten aus dem Katalog. Die gelb hinterlegten Spalten füllt die prüfende Person aus.

    python3 scripts/pruefpaket.py [ziel.xlsx]
"""
import glob
import os
import sys
from datetime import date

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

WURZEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ZIEL = sys.argv[1] if len(sys.argv) > 1 else os.path.join(WURZEL, "Amtlos-Pruefpaket-Pflichten.xlsx")

BEREICH = {
    "gruendung": "Gründung", "gesellschaft": "Gesellschaft", "steuern": "Steuern", "personal": "Personal",
    "arbeitsschutz": "Arbeitsschutz", "datenschutz": "Datenschutz", "it_sicherheit": "IT-Sicherheit",
    "ki": "Künstliche Intelligenz", "online": "Website und Online-Handel", "produkte": "Produkte", "branche": "Branche",
}
PRIO = {"hoch": "A", "mittel": "B", "niedrig": "C"}
ERGEBNISSE = ["korrekt", "anpassen", "falsch", "nicht beurteilt"]

SCHRIFT = "Arial"
GELB = PatternFill("solid", fgColor="FFF4C2")
KOPF = PatternFill("solid", fgColor="063B29")
LINIE = Side(style="thin", color="D5DDD9")


def de(x):
    if isinstance(x, dict):
        return x.get("de", "")
    return "" if x is None else str(x)


def dauer(t):
    """ISO-Dauer (P3D, P1M, P2Y) in Worten."""
    import re
    m = re.fullmatch(r"P(\d+)([DWMY])", str(t or ""))
    if not m:
        return str(t)
    n, e = int(m.group(1)), m.group(2)
    w = {"D": ("Tag", "Tage"), "W": ("Woche", "Wochen"), "M": ("Monat", "Monate"), "Y": ("Jahr", "Jahre")}[e]
    return f"{n} {w[0] if n == 1 else w[1]}"


def wert_text(w):
    if isinstance(w, dict) and "regelwert" in w:
        return f"Regelwert {w['regelwert']}"
    if isinstance(w, bool):
        return "ja" if w else "nein"
    return str(w)


def ausloeser_text(a, tiefe=0):
    """Auslöser-Baum (all/any/not, Bedingungen mit Label) als eingerückter Klartext."""
    pre = "  " * tiefe
    if a is None:
        return ""
    if isinstance(a, dict):
        if "immer" in a:
            return pre + "Gilt immer (für jedes Unternehmen)."
        for schluessel, wort in (("all", "Alle Bedingungen:"), ("any", "Mindestens eine Bedingung:")):
            if schluessel in a:
                teile = [ausloeser_text(x, tiefe + 1) for x in a[schluessel]]
                return "\n".join([pre + wort] + teile) if tiefe or len(teile) > 1 else teile[0].strip()
        if "not" in a:
            return pre + "Nicht:\n" + ausloeser_text(a["not"], tiefe + 1)
        label = de(a.get("label"))
        if "ereignis" in a and "feld" not in a:
            technik = f"[Ereignis {a['ereignis']}]"
        else:
            technik = f"[{a.get('feld', '?')} {a.get('op', '')}" + ("" if a.get("wert") is None else f" {wert_text(a.get('wert'))}") + "]"
        return pre + "• " + (f"{label} {technik}" if label else technik)
    return pre + str(a)


def frist_text(f, w):
    if not f:
        return ""
    typ = f.get("typ")
    teile = []
    if typ == "dauerhaft":
        teile.append("Dauerhaft (ohne Termin)")
    elif typ == "absolut":
        teile.append(f"Stichtag {f.get('datum', '')}".strip())
    elif typ == "relativ":
        bezug = (f.get("bezug") or {}).get("ereignis", "?")
        teile.append(f"Relativ zum Ereignis „{bezug}“, Versatz {dauer(f.get('versatz'))}" + (f", {f['zeitpunkt']}" if f.get("zeitpunkt") else ""))
    elif typ == "wiederkehrend":
        r = f.get("rhythmus")
        if isinstance(r, dict) and "waehle" in r:
            optionen = [f"{o.get('dann')} wenn [{o['wenn'].get('feld')} {o['wenn'].get('op')} {wert_text(o['wenn'].get('wert'))}]" for o in r["waehle"]]
            teile.append("Rhythmus: " + "; ".join(optionen) + f"; sonst {r.get('sonst')}")
        else:
            teile.append(f"Rhythmus: {r}")
        fa = f.get("faellig") or {}
        if fa:
            teile.append("Fällig: " + ", ".join(f"{k} {v}" for k, v in fa.items()))
        if f.get("verlaengerung"):
            teile.append(f"Verlängerung: {dauer(f['verlaengerung'].get('versatz'))} wenn [{f['verlaengerung'].get('wenn', {}).get('feld')}]")
    else:
        teile.append(str(typ))
    v = f.get("verschiebung") or {}
    if v.get("regel") and v.get("regel") != "keine":
        teile.append(f"Wochenende/Feiertag: {v['regel']} ({v.get('kalender', '')})")
    art = (w or {}).get("art")
    if art:
        teile.append({"einmalig": "Einmalig", "periodisch": "Wiederkehrend", "bei_aenderung": "Erneut bei Änderungen"}.get(art, art))
    return "\n".join(teile)


def gueltig(d):
    ab = d.get("gueltig_ab")
    return (f"ab {ab}" if ab else "Beginn offen") + (f" bis {d['gueltig_bis']}" if d.get("gueltig_bis") else "")


def quellen_text(q):
    return "\n".join(f"{x.get('titel', '')}: {x.get('url', '')}" for x in (q or []))


def pflichten():
    zeilen = []
    for pfad in sorted(glob.glob(os.path.join(WURZEL, "catalog", "pflichten", "*.yaml"))):
        d = yaml.safe_load(open(pfad, encoding="utf-8"))
        s = d.get("sanktion") or {}
        zeilen.append({
            "prio": PRIO.get(s.get("schwere"), "B"),
            "nr": d.get("katalog_nr", ""),
            "id": d["id"],
            "bereich": BEREICH.get(d.get("kategorie"), d.get("kategorie", "")),
            "titel": de(d.get("titel")),
            "kurz": de(d.get("kurz")),
            "erklaerung": de(d.get("erklaerung")).strip(),
            "ausloeser": ausloeser_text(d.get("ausloeser")),
            "frist": frist_text(d.get("frist"), d.get("wiederholung")),
            "stelle": de(d.get("stelle")),
            "recht": ", ".join(d.get("rechtsgrundlage") or []),
            "sanktion": f"{de(s.get('text'))} (Einstufung: {s.get('schwere', '?')})",
            "quellen": quellen_text(d.get("quellen")),
            "offen": "\n".join(f"• {x}" for x in (d.get("fachlich_zu_pruefen") or [])),
            "gueltig": gueltig(d),
            "status": f"{d.get('status_recht', '')} · Prüfung: {(d.get('pruefung') or {}).get('status', '')}",
        })
    zeilen.sort(key=lambda z: (z["prio"], z["bereich"], z["titel"]))
    return zeilen


def regelwerte():
    zeilen = []
    for pfad in sorted(glob.glob(os.path.join(WURZEL, "catalog", "regelwerte", "*.yaml"))):
        for d in yaml.safe_load(open(pfad, encoding="utf-8")) or []:
            zeilen.append({
                "nr": d.get("katalog_nr", ""), "id": d["id"], "bezeichnung": de(d.get("bezeichnung")),
                "wert": d.get("wert"), "einheit": d.get("einheit", ""), "gueltig": gueltig(d),
                "recht": ", ".join(d.get("rechtsgrundlage") or []), "quellen": quellen_text(d.get("quellen")),
                "offen": "\n".join(f"• {x}" for x in (d.get("fachlich_zu_pruefen") or [])),
            })
    return zeilen


def tabelle(ws, spalten, zeilen, eingabe_ab):
    """Kopf, Zeilen, Breiten, gelbe Eingabespalten, Auswahl „Ergebnis“, Filter und fixierte Kopfzeile."""
    for i, (titel, _, breite) in enumerate(spalten, 1):
        c = ws.cell(row=1, column=i, value=titel)
        c.font = Font(name=SCHRIFT, bold=True, color="FFFFFF")
        c.fill = KOPF
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = breite
    for r, z in enumerate(zeilen, 2):
        for i, (_, feld, _) in enumerate(spalten, 1):
            c = ws.cell(row=r, column=i, value=z.get(feld) if feld else None)
            c.font = Font(name=SCHRIFT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=LINIE)
            if i >= eingabe_ab:
                c.fill = GELB
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "F2" if ws.title == "Pflichten" else "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spalten))}{len(zeilen) + 1}"
    dv = DataValidation(type="list", formula1='"' + ",".join(ERGEBNISSE) + '"', allow_blank=True)
    dv.error = "Bitte einen Wert aus der Liste wählen."
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(eingabe_ab)}2:{get_column_letter(eingabe_ab)}{len(zeilen) + 1}")


def main():
    p = pflichten()
    rw = regelwerte()
    wb = Workbook()

    an = wb.active
    an.title = "Anleitung"
    an.column_dimensions["A"].width = 34
    an.column_dimensions["B"].width = 90
    zeilen = [
        ("Amtlos – Prüfpaket Pflichtenkatalog", None),
        (f"Stand {date.today().strftime('%d.%m.%Y')}, erzeugt aus dem Katalog im Code (scripts/pruefpaket.py).", None),
        (None, None),
        ("Worum es geht", "Amtlos zeigt Selbstständigen, welche gesetzlichen Pflichten für ihren Betrieb gelten und wann. Alle Einträge sind Entwürfe und noch nicht fachlich freigegeben. Bitte prüfen Sie je Pflicht, ob Auslöser, Frist, Rechtsgrundlage und Beschreibung zutreffen."),
        ("Was auszufüllen ist", "Nur die gelb hinterlegten Spalten am rechten Rand: „Ergebnis“ (Auswahl), „Korrektur / Anmerkung“, „Geprüft von“, „Datum“. Alles andere kommt aus dem Katalog."),
        ("Ergebnis – Bedeutung", "korrekt = so freigeben · anpassen = im Kern richtig, Korrektur in der Anmerkung · falsch = Pflicht/Auslöser/Frist trifft nicht zu · nicht beurteilt = außerhalb Ihres Fachgebiets"),
        ("Beispiel für eine ausgefüllte Zeile", "Ergebnis: anpassen · Anmerkung: „Frist gilt erst ab Einstellung der ersten Person mit Arbeitsvertrag, nicht bei Minijob.“ · Geprüft von: RA Mustermann · Datum: 20.10.2026"),
        ("Priorität", "A = Sanktion laut Katalog „hoch“ (zuerst prüfen) · B = mittel · C = niedrig"),
        ("Auslöser lesen", "In Klartext die Bedingungen aus dem Fragebogen der App; in eckigen Klammern das technische Feld, z. B. [profil.beschaeftigte.gesamt >= 1]. „Regelwert …“ verweist auf das Blatt „Regelwerte“ (Beträge und Grenzen)."),
        ("Offene Prüfpunkte", "Fragen, die beim Erfassen schon aufgefallen sind – bitte besonders darauf achten."),
        (None, None),
        ("Fortschritt", None),
    ]
    for r, (a, b) in enumerate(zeilen, 1):
        if a is not None:
            an.cell(row=r, column=1, value=a).font = Font(name=SCHRIFT, bold=(r == 1 or b is not None or a == "Fortschritt"), size=16 if r == 1 else 11)
        if b is not None:
            c = an.cell(row=r, column=2, value=b)
            c.font = Font(name=SCHRIFT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            an.row_dimensions[r].height = 46
    start = len(zeilen) + 1
    n = len(p) + 1
    for i, erg in enumerate(ERGEBNISSE + ["offen"], 0):
        an.cell(row=start + i, column=1, value=f"Pflichten {erg}").font = Font(name=SCHRIFT, size=11)
        formel = f'=COUNTBLANK(Pflichten!Q2:Q{n})' if erg == "offen" else f'=COUNTIF(Pflichten!Q2:Q{n},"{erg}")'
        an.cell(row=start + i, column=2, value=formel).font = Font(name=SCHRIFT, size=11)
    an.cell(row=start + 5, column=1, value="Pflichten gesamt").font = Font(name=SCHRIFT, bold=True)
    an.cell(row=start + 5, column=2, value=f"=COUNTA(Pflichten!C2:C{n})").font = Font(name=SCHRIFT, bold=True)
    m = len(rw) + 1
    an.cell(row=start + 6, column=1, value="Regelwerte geprüft").font = Font(name=SCHRIFT, size=11)
    an.cell(row=start + 6, column=2, value=f'=COUNTA(Regelwerte!J2:J{m})&" von "&COUNTA(Regelwerte!B2:B{m})').font = Font(name=SCHRIFT, size=11)
    for r in range(start, start + 7):
        an.cell(row=r, column=2).alignment = Alignment(horizontal="left")

    ws = wb.create_sheet("Pflichten")
    tabelle(ws, [
        ("Prio", "prio", 6), ("Katalog-Nr.", "nr", 10), ("ID", "id", 28), ("Bereich", "bereich", 16), ("Pflicht", "titel", 30),
        ("Beschreibung in der App", "kurz", 44), ("Erklärung in der App", "erklaerung", 50), ("Gilt, wenn …", "ausloeser", 52),
        ("Frist / Wiederholung", "frist", 40), ("Zuständige Stelle", "stelle", 28), ("Rechtsgrundlage", "recht", 22),
        ("Sanktion", "sanktion", 32), ("Quellen", "quellen", 40), ("Offene Prüfpunkte", "offen", 46), ("Gültig", "gueltig", 16),
        ("Status", "status", 20), ("Ergebnis", None, 14), ("Korrektur / Anmerkung", None, 50), ("Geprüft von", None, 18), ("Datum", None, 12),
    ], p, 17)

    ws2 = wb.create_sheet("Regelwerte")
    tabelle(ws2, [
        ("Katalog-Nr.", "nr", 10), ("ID", "id", 34), ("Bezeichnung", "bezeichnung", 48), ("Wert", "wert", 12), ("Einheit", "einheit", 9),
        ("Gültig", "gueltig", 18), ("Rechtsgrundlage", "recht", 22), ("Quellen", "quellen", 40), ("Offene Prüfpunkte", "offen", 46),
        ("Ergebnis", None, 14), ("Korrektur / Anmerkung", None, 50), ("Geprüft von", None, 18), ("Datum", None, 12),
    ], rw, 10)

    wb.save(ZIEL)
    print(f"{len(p)} Pflichten, {len(rw)} Regelwerte → {ZIEL}")


if __name__ == "__main__":
    main()
