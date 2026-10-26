#!/usr/bin/env python3
"""
E-211: Prüfpakete für die Regelwerte (catalog/regelwerte – Schwellen, Beträge, Sätze) – eine Excel-Datei je Rolle
(Steuerberater, Rechtsanwalt, Fachkraft für Arbeitssicherheit), im Stil von scripts/pruefpaket.py. Je Wert: Bezeichnung,
Wert mit Einheit, Gültigkeit, Rechtsstand, Rechtsgrundlage, Quelle, offene Prüfpunkte und wo er verwendet wird (Auslöser,
Frist, Text einer Pflicht; Code). Gelb: „Ergebnis (richtig / falsch → richtiger Wert)“ und „Anmerkung“. Werte werden nicht
verändert. Rücklauf: docs/freigabe-katalog/regelwerte-uebernehmen.py.

    python3 scripts/regelwerte-pruefpaket.py <ordner>     # außerhalb des Repositorys (E-199)
"""
import os
import sys
from datetime import date

WURZEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(WURZEL, "docs/freigabe-katalog"))
from ziel import ziel_ausserhalb  # noqa: E402
import regelwerte as rw  # noqa: E402

ZIEL = ziel_ausserhalb(sys.argv[1] if len(sys.argv) > 1 else None, "Aufruf: python3 scripts/regelwerte-pruefpaket.py <ordner>")
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

SCHRIFT = "Arial"
GELB = PatternFill("solid", fgColor="FFF4C2")
KOPF = PatternFill("solid", fgColor="063B29")
LINIE = Side(style="thin", color="D5DDD9")
DATEI = {"steuern": "Regelwerte-Steuerberater.xlsx", "recht": "Regelwerte-Rechtsanwalt.xlsx", "arbeitsschutz": "Regelwerte-Fachkraft-Arbeitssicherheit.xlsx"}
RECHTSSTAND = {"in_kraft": "in Kraft", "beschlossen": "beschlossen, noch nicht in Kraft", "entwurf": "Entwurf", "geplant": "geplant", "ausgelaufen": "ausgelaufen"}
ERGEBNIS = "Ergebnis (richtig / falsch → richtiger Wert)"
SPALTEN = [
    ("Nr.", "nr", 5), ("Katalog-Nr.", "katalog_nr", 9), ("ID", "id", 34), ("Version", "version", 8), ("Bezeichnung", "bezeichnung", 44),
    ("Wert + Einheit", "wert", 16), ("gültig ab", "ab", 11), ("gültig bis", "bis", 11), ("Rechtsstand", "rechtsstand", 14),
    ("Rechtsgrundlage", "recht", 22), ("Quelle", "quelle", 40), ("offene Prüfpunkte", "offen", 46), ("verwendet in", "verwendet", 52),
    (ERGEBNIS, None, 24), ("Anmerkung", None, 44),
]
EINGABE_AB = 14


def datum(x):
    return "" if not x else ".".join(reversed(str(x).split("-")))


def zeilen_je_rolle():
    nutzung = rw.verwendung(rw.pflichten_laden())
    je = {r: [] for r in rw.ROLLEN}
    for _, d in rw.laden():
        verwendet = [f"{art}: {titel} ({pid})" for art, pid, titel in nutzung.get(d["id"], [])]
        if d["id"] in rw.CODE_VERWENDUNG:
            verwendet.append(rw.CODE_VERWENDUNG[d["id"]])
        je[rw.rolle(d["id"])].append({
            "katalog_nr": d.get("katalog_nr", ""), "id": d["id"], "version": d["version"], "bezeichnung": rw.de(d.get("bezeichnung")),
            "wert": rw.wert_text(d["wert"], d.get("einheit", "")), "ab": datum(d.get("gueltig_ab")), "bis": datum(d.get("gueltig_bis")) or "offen",
            "rechtsstand": RECHTSSTAND.get(d.get("status_recht"), d.get("status_recht", "")), "recht": ", ".join(d.get("rechtsgrundlage") or []),
            "quelle": "\n".join(f"{q.get('titel', '')}: {q.get('url', '')}" for q in d.get("quellen") or []),
            "offen": "\n".join(f"• {x}" for x in d.get("fachlich_zu_pruefen") or []),
            "verwendet": "\n".join(verwendet) or "kein Verweis aus einer Pflicht oder dem Code gefunden (steht der Wert sinngemäß in einem Pflichttext, bitte in der Anmerkung nennen)",
        })
    for r in je:
        for i, z in enumerate(je[r], 1):
            z["nr"] = i
    return je


def mappe(rolle, zeilen, pfad):
    wb = Workbook()
    an = wb.active
    an.title = "Anleitung"
    an.column_dimensions["A"].width = 30
    an.column_dimensions["B"].width = 100
    text = [
        (f"Prüfliste Regelwerte – {rw.ROLLEN[rolle]}", None),
        (f"Stand {date.today().strftime('%d.%m.%Y')}, erzeugt aus dem Katalog im Code (scripts/regelwerte-pruefpaket.py).", None),
        (None, None),
        ("Worum es geht", "Die App entscheidet mit diesen Werten (Schwellen, Beträge, Sätze), ob eine Pflicht für einen Betrieb gilt und in welchem Rhythmus. "
         "Die Pflichten selbst sind fachlich geprüft; diese Werte wurden dabei nur als Verweis gezeigt und sind noch nicht bestätigt. "
         "Die App sagt das bei jeder betroffenen Pflicht („Schwellenwert … noch nicht fachlich bestätigt“), bis der Wert bestätigt ist."),
        ("Was auszufüllen ist", f"Nur die gelben Spalten: „{ERGEBNIS}“ – „richtig“ oder „falsch → richtiger Wert“ (z. B. „falsch → 22.000 €“) – und „Anmerkung“ "
         "(z. B. zur Grenze inklusive/exklusive, zum Bezugsjahr oder zu „gültig ab“). Bitte keine anderen Zellen ändern."),
        ("Offene Prüfpunkte", "Fragen, die beim Erfassen aufgefallen sind (oft: „gültig ab“ ist ein Platzhalter, Vergleich > oder ≥) – bitte in der Anmerkung beantworten."),
        ("Verwendet in", "Auslöser = der Wert entscheidet, ob die Pflicht gilt · Frist = der Wert bestimmt den Rhythmus · Text = der Betrag steht im Text der Pflicht · "
         "App/Server = der Wert wird außerhalb einer Pflicht angezeigt oder berechnet."),
        ("Bestätigung", "Bitte die ausgefüllte Datei zurückschicken und in der Mail bestätigen, dass die mit „richtig“ markierten Werte zutreffen. "
         "Ein Name erscheint nie in der App; dort steht nur die Rolle."),
        (None, None),
        ("Werte in dieser Liste", len(zeilen)),
    ]
    for r, (a, b) in enumerate(text, 1):
        if a is not None:
            an.cell(row=r, column=1, value=a).font = Font(name=SCHRIFT, bold=(r == 1 or b is not None), size=16 if r == 1 else 11)
        if b is not None:
            c = an.cell(row=r, column=2, value=b)
            c.font = Font(name=SCHRIFT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            an.row_dimensions[r].height = 48 if isinstance(b, str) else 16
    ws = wb.create_sheet("Regelwerte")
    for i, (titel, _, breite) in enumerate(SPALTEN, 1):
        c = ws.cell(row=1, column=i, value=titel)
        c.font = Font(name=SCHRIFT, bold=True, color="FFFFFF")
        c.fill = KOPF
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = breite
    for r, z in enumerate(zeilen, 2):
        for i, (_, feld, _) in enumerate(SPALTEN, 1):
            c = ws.cell(row=r, column=i, value=z.get(feld) if feld else None)
            c.font = Font(name=SCHRIFT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=LINIE)
            if i >= EINGABE_AB:
                c.fill = GELB
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{len(zeilen) + 1}"
    wb.save(pfad)


def main():
    os.makedirs(ZIEL, exist_ok=True)
    for rolle, zeilen in zeilen_je_rolle().items():
        if not zeilen:
            continue
        pfad = os.path.join(ZIEL, DATEI[rolle])
        mappe(rolle, zeilen, pfad)
        print(f"{rw.ROLLEN[rolle]}: {len(zeilen)} Regelwerte → {pfad}")


if __name__ == "__main__":
    main()
